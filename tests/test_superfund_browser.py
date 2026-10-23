"""Actual browser layout checks; enabled explicitly by the release gate."""
import os
import socket
from pathlib import Path
import subprocess
import sys
import time
from urllib.request import urlopen
import pytest

pytestmark=pytest.mark.skipif(os.getenv('SF_BROWSER_TEST')!='1',reason='Browser checks run in release CI')


@pytest.fixture(scope='module')
def server():
    env={**os.environ,'PYTHONPATH':'.'}
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
    base='http://127.0.0.1:'+str(port)
    process=subprocess.Popen([sys.executable,'-m','streamlit','run','scripts/superfund_ui_fixture.py','--server.port='+str(port),'--server.headless=true','--browser.gatherUsageStats=false','--theme.base=dark'],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    try:
        for _ in range(100):
            try:
                if urlopen(base+'/_stcore/health',timeout=1).status==200:break
            except Exception:time.sleep(.2)
        else:raise AssertionError('Fixture server did not start')
        yield base
    finally:process.terminate();process.wait(timeout=10)


@pytest.mark.parametrize('width',[375,430,1280])
def test_superfund_all_menus_wrap_without_overlap(server,width):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser=pw.chromium.launch(executable_path=os.getenv("BROWSER_EXECUTABLE") or None)
        p=browser.new_page(viewport={'width':width,'height':900},device_scale_factor=1)
        p.goto(server+'?aa_nav=superfund')
        p.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        p.locator('.sf-card').first.wait_for()
        Path('dist/browser').mkdir(parents=True,exist_ok=True)
        p.locator('.sf-card').first.scroll_into_view_if_needed()
        p.screenshot(path=f'dist/browser/superfund-pending-{width}.png')
        # Navigate using the app's actual section selector. Each selection
        # triggers a Streamlit rerun; do not rely on every expander staying open.
        for label in ['Kandidater','Rapporter og nedlastinger','Nyheter og kilder','Læring og skygge','Parametre','Diagnoser']:
            selector=p.get_by_role('combobox',name='Vis område')
            selector.click()
            p.get_by_role('option',name=label,exact=True).click()
            p.locator('[data-testid="stExpander"] summary').filter(has_text=label).first.wait_for()
            assert p.locator('[data-testid="stExpander"] summary').filter(has_text=label).count()==1
            assert p.get_by_text('Uncaught app exception',exact=False).count()==0
            if label=='Diagnoser':
                p.locator('.sf-card').get_by_text('1725.7',exact=True).first.wait_for()
                p.screenshot(path=f'dist/browser/superfund-diagnoses-{width}.png')
        # Restore the candidate panel for the ranking assertions below.
        p.get_by_role('combobox',name='Vis område').click()
        p.get_by_role('option',name='Kandidater',exact=True).click()
        p.get_by_role('heading',name='Topp 25 · grunnkrav oppfylt',exact=True).wait_for()
        problems=p.evaluate('''() => {
          const errors=[];
          for (const card of document.querySelectorAll('.sf-card')) {
            if (!card.getClientRects().length) continue;
            const box=card.getBoundingClientRect();
            if (box.right>innerWidth+2 || box.left< -2) errors.push('card outside viewport');
            let bottom=card.querySelector('h3').getBoundingClientRect().bottom;
            for(const fact of card.querySelectorAll('.sf-fact')) {
              const r=fact.getBoundingClientRect();
              if(r.top<bottom-1) errors.push('overlapping facts');
              if(r.bottom>box.bottom+1) errors.push('fact outside card');
              bottom=r.bottom;
            }
          }
          for (const caption of document.querySelectorAll('.st-key-superfund_page [data-testid="stCaptionContainer"]')) {
            if(caption.getClientRects().length && caption.scrollHeight>caption.clientHeight+2) errors.push('caption clipped');
            if(caption.getClientRects().length && parseFloat(getComputedStyle(caption).fontSize)<14) errors.push('caption too small');
          }
          for (const a of document.querySelectorAll('.st-key-superfund_page [data-testid="stLinkButton"] a')) {
            if(!a.getClientRects().length)continue;
            const style=getComputedStyle(a);
            if(style.color===style.backgroundColor) errors.push('invisible link');
          }
          return errors;
        }''')
        Path('dist/browser').mkdir(parents=True,exist_ok=True)
        p.get_by_role('heading',name='Topp 25 · grunnkrav oppfylt',exact=True).scroll_into_view_if_needed()
        p.screenshot(path=f'dist/browser/superfund-top25-{width}.png')
        assert not problems,problems
        browser.close()


@pytest.mark.parametrize('width',[375,430])
def test_first_tap_and_fresh_browser_restore_superfund_and_selected_section(server,width):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser=pw.chromium.launch(executable_path=os.getenv("BROWSER_EXECUTABLE") or None)
        p=browser.new_page(viewport={'width':width,'height':900})
        p.goto(server+'?aa_nav=dashboard')
        p.get_by_role('heading',name='Start',exact=True).wait_for()
        p.locator('.st-key-aa_mobile_nav_superfund button').click()
        p.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        assert 'aa_nav=superfund' in p.url
        # Changing the selected section updates the durable route.
        p.get_by_role('combobox',name='Vis område').click()
        p.get_by_role('option',name='Kandidater',exact=True).click()
        p.wait_for_function("() => decodeURIComponent(location.search).includes('aa_tab=Kandidater')")
        p.reload()
        p.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        p.locator('[data-testid="stExpander"] summary').filter(has_text='Kandidater').first.wait_for()
        p.get_by_role('combobox',name='Vis område').wait_for()
        # Streamlit can complete the selector render before candidate cards finish rerunning.
        # Wait for the actual candidate tag, rather than checking visibility immediately.
        p.locator('.sf-tag').first.wait_for(state='visible', timeout=15000)
        # Pending model orders are shown in Portefølje, not Kandidater.
        p.get_by_role('combobox',name='Vis område').click()
        p.get_by_role('option',name='Portefølje',exact=True).click()
        p.locator('[data-testid="stExpander"] summary').filter(has_text='Portefølje').first.wait_for()
        p.locator('.sf-card').filter(has_text='Kjøp avventer ny kursobservasjon').first.wait_for()
        p.locator('.st-key-aa_mobile_nav_overview button').click()
        p.get_by_role('heading',name='Start',exact=True).wait_for()
        p.locator('.st-key-aa_mobile_nav_superfund button').click()
        p.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        # Verify persisted Portfolio selection after returning through the mobile nav.
        p.locator('[data-testid="stExpander"] summary').filter(has_text='Portefølje').first.wait_for()
        p.locator('.sf-card').filter(has_text='Kjøp avventer ny kursobservasjon').first.wait_for()
        # Candidate tags must only be asserted when Candidate view is selected.
        p.get_by_role('combobox',name='Vis område').click()
        p.get_by_role('option',name='Kandidater',exact=True).click()
        p.locator('.sf-tag').first.wait_for(state='visible')
        Path('dist/browser').mkdir(parents=True,exist_ok=True)
        p.screenshot(path=f'dist/browser/superfund-first-tap-{width}.png')
        browser.close()


@pytest.mark.parametrize('width',[375,430])
def test_country_ranking_switch_shows_separate_brazil_group(server,width):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser=pw.chromium.launch(executable_path=os.getenv("BROWSER_EXECUTABLE") or None)
        p=browser.new_page(viewport={'width':width,'height':900})
        p.goto(server+'?aa_nav=superfund&aa_tab=Kandidater')
        p.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        p.get_by_text('Beste per investeringsland/region',exact=True).click()
        group=p.locator('[data-testid="stExpander"] summary').filter(has_text='Brasil ·')
        group.first.wait_for();group.first.click()
        group.first.locator('xpath=..').get_by_text('Plass 1 innen området',exact=True).wait_for()
        assert p.locator('.sf-card').filter(has_text='Siste uke').count()>0
        group.first.scroll_into_view_if_needed()
        Path('dist/browser').mkdir(parents=True,exist_ok=True)
        p.screenshot(path=f'dist/browser/superfund-country-{width}.png')
        browser.close()


@pytest.mark.parametrize('width',[375,1280])
def test_superfund_downloaded_pdf_csv_excel_and_diagnostic_zip(server,width):
    from io import BytesIO
    import csv
    import zipfile
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        browser=pw.chromium.launch(executable_path=os.getenv('BROWSER_EXECUTABLE') or None)
        page=browser.new_page(viewport={'width':width,'height':900},accept_downloads=True)
        page.goto(server+'?aa_nav=superfund')
        page.get_by_role('heading',name='Superfondportefølje',exact=True).wait_for()
        page.get_by_role('combobox',name='Vis område').click()
        page.get_by_role('option',name='Rapporter og nedlastinger',exact=True).click()
        page.get_by_role('button',name='Lag rapporter for nedlasting',exact=True).click()
        files={}
        for kind in ['PDF','CSV','Excel']:
            with page.expect_download() as download:
                page.get_by_role('button',name='Last ned '+kind,exact=True).click()
            files[kind]=Path(download.value.path()).read_bytes()
        assert len(PdfReader(BytesIO(files['PDF'])).pages)>0
        assert list(csv.reader(files['CSV'].decode('utf-8-sig').splitlines()))
        assert load_workbook(BytesIO(files['Excel']),read_only=True).sheetnames
        # Diagnose-ZIP belongs to Rapporter og nedlastinger, not Diagnoser.
        page.get_by_role('button',name='Lag diagnose-ZIP',exact=True).click()
        with page.expect_download() as download:
            page.get_by_role('button',name='Last ned diagnose-ZIP',exact=True).click()
        with zipfile.ZipFile(download.value.path()) as archive:
            assert archive.testzip() is None and archive.namelist()
        assert not page.locator('[data-testid="stException"]').count()
        browser.close()
