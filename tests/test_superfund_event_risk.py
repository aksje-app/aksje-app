from copy import deepcopy
import pytest
import superfund_engine as engine
import superfund_event_risk as risk
import superfund_runtime as runtime
from test_superfund_v1 import item,step,local


def br(key='NO0000000001:78',day='2026-10-12',**kw):
    row=item(key,day=day,**kw)
    row['detail']['regions']=[{'exposureType':'BRAZIL','weight':100}]
    return row


def usa(key='NO0000000002:78',day='2026-10-12',**kw):
    row=item(key,day=day,**kw)
    row['detail']['regions']=[{'exposureType':'USA','weight':100}]
    return row


def held(row,allocation=200000):
    state,_=step(None,[row],day='2026-10-09')
    # Seed a legitimate legacy holding above the new event cap.
    state['orders']=[{'id':row['id'],'side':'BUY','allocation':allocation,'requested_at':'2026-10-08T15:00:00+00:00','next_at':'2026-10-09T00:00:00+00:00'}]
    state,_=step(state,[row],day='2026-10-09',p={'max_position_pct':25})
    return state


def test_stop_clamps_legacy_and_shadow_parameters():
    assert engine.valid_parameters({**engine.DEFAULTS,'stop_pct':6})['stop_pct']==3
    s=held(br(day='2026-10-09'))
    s,_=step(s,[br(day='2026-10-12',price=96)],'2026-10-12',{'stop_pct':6})
    assert s['orders'][0]['side']=='SELL'


def test_cap_includes_pending_orders_and_same_currency_different_funds():
    rows=[br(),br('NO0000000003:78',category='Other'),usa()]
    s,r=step(None,rows,'2026-10-12')
    assert sum(o['allocation'] for o in s['orders'] if o['isin']!='NO0000000002')<=100000
    assert any(o['isin']=='NO0000000002' for o in s['orders'])
    later=[br(day='2026-10-13'),br('NO0000000003:78',day='2026-10-13',category='Other'),usa(day='2026-10-13')]
    s,_=step(s,later,'2026-10-13')
    equity=s['cash']+sum(x['quantity']*x['last_nok'] for x in s['positions'].values())
    assert risk.exposure_used(s['positions'],[],{})<=equity*.1+.01


def test_changed_exposure_at_fill_is_rechecked():
    s,_=step(None,[usa(),usa('NO0000000003:78',category='Other')],'2026-10-12')
    s,_=step(s,[br('NO0000000002:78',day='2026-10-13'),br('NO0000000003:78',day='2026-10-13',category='Other')],'2026-10-13')
    assert sum(x['quantity']*x['last_nok'] for x in s['positions'].values())<=100000


def test_no_political_assumptions_or_isin_country_inference():
    row=usa('BR0000000001:78')
    assert risk.brazil_weight(row)[0]==0
    row=br();row['detail']['regions']=[]
    assert risk.brazil_weight(row)[0]==1
    row=br();row['detail']['regions']=[{'exposureType':'BRAZIL','weight':30},{'exposureType':'USA','weight':70}]
    assert risk.brazil_weight(row)[0]==pytest.approx(.3)


def test_proposal_requires_approval_partial_accounting_and_no_replay():
    s=held(br(day='2026-10-09'))
    s,_=step(s,[br()],'2026-10-12')
    assert not s['orders']
    view=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)
    proposal=view['proposals'][0];old=deepcopy(s)
    approvals={proposal['id']:{'state':'APPROVED'}}
    s,_=engine.cycle(s,[br()],'2026-10-12T15:00:00+00:00',approvals=approvals)
    assert len(s['orders'])==1 and s['orders'][0]['fraction']<1
    s,_=engine.cycle(s,[br(day='2026-10-13')],'2026-10-13T15:00:00+00:00',approvals=approvals)
    pos=s['positions'][br()['id']]
    assert pos['quantity']<old['positions'][br()['id']]['quantity']
    assert pos['cost_nok']/old['positions'][br()['id']]['cost_nok']==pytest.approx(pos['quantity']/old['positions'][br()['id']]['quantity'])
    repeat,_=engine.cycle(s,[br(day='2026-10-13')],'2026-10-13T15:00:00+00:00',approvals=approvals)
    assert repeat==s
    assert len(s['trades'])==2


def test_stale_approval_does_not_sell_changed_position():
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    p=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)['proposals'][0]
    s,_=engine.cycle(s,[br(day='2026-10-13',price=101)],'2026-10-13T15:00:00+00:00',approvals={p['id']:{'state':'APPROVED'}})
    assert not s['orders']


def test_event_quarantine_covers_alternate_fund_and_two_new_dates():
    s=engine.initial();risk.register_exit(s,br(),engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)
    alternate=br('NO0000000003:78',day='2026-10-25')
    assert 'HENDELSESKARANTENE' in risk.reentry_blocks(s,alternate,engine.at('2026-10-25T15:00:00+00:00'))
    assert not risk.reentry_blocks(s,usa(),engine.at('2026-10-25T15:00:00+00:00'))
    s,_=step(s,[br('NO0000000003:78',day='2026-10-26')],'2026-10-26')
    assert not s['orders']
    s,_=engine.cycle(s,[br('NO0000000003:78',day='2026-10-26')],'2026-10-26T16:00:00+00:00')
    assert not s['orders']
    s,_=step(s,[br('NO0000000003:78',day='2026-10-27')],'2026-10-27')
    assert not s['orders']  # midnight UTC NAV predates the Brazil-local expiry on the 26th
    s,_=step(s,[br('NO0000000003:78',day='2026-10-28')],'2026-10-28')
    assert s['orders'][0]['side']=='BUY'


def test_pending_risk_exit_blocks_replacement_before_fill():
    s=held(br(day='2026-10-09'));rows=[br(price=95),br('NO0000000003:78')]
    s,r=step(s,rows,'2026-10-12')
    assert [o['side'] for o in s['orders']]==['SELL']


def test_time_window_rechecks_same_quotes_without_fake_confirmations():
    assert not risk.active(engine.at('2026-10-11T01:00:00+00:00'))
    assert risk.active(engine.at('2026-10-11T04:00:00+00:00'))
    assert not risk.active(engine.at('2026-10-26T03:00:00+00:00'))


def test_cached_holdings_overlap_blocks_pending_replacement():
    a,b=usa(),usa('NO0000000003:78',category='Other')
    for row in (a,b):row['detail']['holdings']=[{'name':'Same company','weight':60}]
    s,_=step(None,[a,b],'2026-10-12')
    assert len(s['orders'])==1


def test_risk_decision_atomic_receipt_preserves_portfolio(local):
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    view=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)
    runtime.write('snapshot.json',{'model':s,'risk':view})
    pid=view['proposals'][0]['id']
    assert runtime.decide_risk_proposal(pid,True)['state']=='APPROVED'
    assert runtime.decide_risk_proposal(pid,False)['state']=='APPROVED'
    assert runtime.snapshot()['model']==s
    with pytest.raises(ValueError):runtime.decide_risk_proposal('obsolete',True)


def test_risk_ui_reads_cache_and_decides(local):
    from streamlit.testing.v1 import AppTest
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    view=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)
    runtime.write('snapshot.json',{'model':s,'risk':view,'parameters':engine.DEFAULTS,'at':'2026-10-12'})
    app=AppTest.from_string('import streamlit as st\nfrom pages.superfund import render_superfund\nrender_superfund(st)').run()
    assert not app.exception
    # Legacy model decisions are not offered on new-period portfolio screens.
    pid=view['proposals'][0]['id']
    assert not [b for b in app.button if b.key=='sf_risk_yes_'+pid]
    assert not runtime.read('risk_decisions.json',{}).get(pid)



def test_receipt_survives_new_quote_but_cannot_repeat():
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    proposal=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)['proposals'][0]
    approvals={proposal['id']:{'state':'APPROVED','at':'2026-10-12T16:00:00+00:00','proposal':proposal}}
    s,_=engine.cycle(s,[br(day='2026-10-13',price=101)],'2026-10-13T15:00:00+00:00',approvals=approvals)
    assert s['trades'][-1]['approval_id']==proposal['id']
    assert s['trades'][-1]['side']=='SELL'
    s,_=engine.cycle(s,[br(day='2026-10-14',price=101)],'2026-10-14T15:00:00+00:00',approvals=approvals)
    assert sum(x.get('approval_id')==proposal['id'] for x in s['trades'])==1


def test_soft_trend_needs_two_quotes_hard_stop_immediate():
    s=held(usa(day='2026-10-09'))
    s,_=step(s,[usa(day='2026-10-12',week=-2)],'2026-10-12')
    assert not s['orders']
    s,_=engine.cycle(s,[usa(day='2026-10-12',week=-2)],'2026-10-12T16:00:00+00:00')
    assert not s['orders']
    s,_=step(s,[usa(day='2026-10-13',week=-2)],'2026-10-13')
    assert s['orders'][0]['side']=='SELL'


def test_approval_cannot_apply_to_replaced_quantity():
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    proposal=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)['proposals'][0]
    s['positions'][br()['id']]['quantity']+=10
    s,_=engine.cycle(s,[br(day='2026-10-13')],'2026-10-13T15:00:00+00:00',approvals={proposal['id']:{'state':'APPROVED','proposal':proposal}})
    assert not s['orders']


def test_exit_execution_blocks_alternative_in_same_cycle():
    s=held(br(day='2026-10-09'))
    s,_=step(s,[br(price=95)],'2026-10-12')
    s,r=step(s,[br(day='2026-10-13',price=95),br('NO0000000003:78',day='2026-10-13')],'2026-10-13')
    assert not s['positions'] and not s['orders']
    assert 'HENDELSESKARANTENE' in r['rows'][1]['blocks']


def test_event_exit_delayed_until_after_election_still_quarantines():
    s=held(br(day='2026-10-09'))
    # Existing decision in event window, late fund settlement after event.
    s['orders']=[{'id':br()['id'],'side':'SELL','requested_at':'2026-10-23T15:00:00+00:00','next_at':'2026-10-26T12:00:00+00:00','reason':'Risiko'}]
    s,_=step(s,[br(day='2026-10-26')],'2026-10-26')
    assert engine.at(s['event_exits'][br()['isin']]['until'])>=engine.at('2026-11-02T15:00:00+00:00')


def test_event_activation_same_utc_day_rechecks_pending_allocation():
    row=br(day='2026-10-10')
    s,_=engine.cycle(None,[row],'2026-10-11T01:00:00+00:00')
    assert s['orders'][0]['allocation']==125000
    s,_=engine.cycle(s,[row],'2026-10-11T04:00:00+00:00')
    assert s['orders'][0]['allocation']<=100000


def test_exports_include_risk_scenario_and_proposals():
    from io import BytesIO
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from superfund_reports import xlsx_bytes,pdf_bytes
    s=held(br(day='2026-10-09'));s,_=step(s,[br()],'2026-10-12')
    view=risk.risk_view(s,[br()],engine.at('2026-10-12T15:00:00+00:00'),engine.DEFAULTS)
    snapshot={'model':s,'risk':view,'parameters':engine.DEFAULTS,'at':'2026-10-12'}
    wb=load_workbook(BytesIO(xlsx_bytes(snapshot)))
    assert 'Hendelsesrisiko' in wb.sheetnames
    text=' '.join(page.extract_text() for page in PdfReader(BytesIO(pdf_bytes(snapshot))).pages)
    assert 'Krever godkjenning' in text and 'Ikke prognose' in text


def test_policy_upgrade_starts_new_learning_and_publishes_effective_shadows(local,monkeypatch):
    from test_superfund_v1 import raw_row
    runtime.write('snapshot.json',{'learning':{'paired_dates':80,'status':'old'},'shadow_rules':{'fast':{**engine.DEFAULTS,'stop_pct':6}}})
    provider=lambda kind,page:{'kind':kind,'page':page,'total':1,'rows':[raw_row('NO0000000001')],'url':'fixture'}
    result=runtime.run_batch(2,page_provider=provider,fx_provider=lambda:{'days':{}},enrich=False)
    assert result['state']=='COMPLETED'
    snapshot=runtime.snapshot()
    assert snapshot['prior_learning']['paired_dates']==80
    assert snapshot['learning']['paired_dates']==0
    assert snapshot['shadow_rules']['fast']['stop_pct']==3
    assert snapshot['risk_policy_version']==1


def test_known_brazil_plus_broad_region_uses_conservative_upper_bound():
    row=br()
    row['detail']['regions']=[{'exposureType':'BRAZIL','weight':30},{'exposureType':'GLOBAL','weight':70}]
    assert risk.brazil_weight(row)[0]==1
    assert 'øvre grense' in risk.brazil_weight(row)[1]
