from io import BytesIO
import csv
from superfund_reports import csv_bytes, xlsx_bytes

def test_export_live_score_columns_and_missing_evidence():
    snapshot={"model":{"positions":{"a":{"name":"Fund A","isin":"AA","category":"Equity","quantity":1,
        "last_nok":100,"cost_nok":100,"bought_at":"2026-10-13T12:00:00Z",
        "rank_at_buy":1,"rank_now":2,"score_at_buy":5}}},
        "candidates":[{"id":"a","returns":{"yield_1w":2,"yield_1m":4,"yield_3m":6},
        "price_at":"2026-10-14T12:00:00Z"}]}
    rows=list(csv.reader(csv_bytes(snapshot).decode("utf-8-sig").splitlines()))
    assert rows[0][6:9]==["Poeng ved kjøp","Poeng nå","Rang blant eide"]
    assert float(rows[1][7])==3.1
    assert rows[1][8]=="1"
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(xlsx_bytes(snapshot)))
    ws=wb["Superfond"]
    assert ws.cell(2,8).value==3.1
    assert ws.cell(2,9).value==1
