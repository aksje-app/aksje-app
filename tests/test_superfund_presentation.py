from copy import deepcopy
from html import unescape
from io import BytesIO
import pytest
from superfund_presentation import candidate_views, top_candidates, candidate_card, card, identity_lines, resolve_identity, date_label


def row(n, **extra):
    return {'id':f'NO{n:010d}:78','isin':f'NO{n:010d}','name':f'Fond {n}', 'kind':'fond',
            'category':'Global Equity','returns_currency':'NOK','score':1000-n,'rank':n+1,'blocks':[],
            'returns':{'yield_1w':2,'yield_1m':3,'yield_3m':4},'risk':4,'cost_pct':.3,**extra}


def test_full_universe_selection_passes_first_hundred_blocked():
    rows=[row(n,blocks=['RISIKO']) for n in range(110)]+[row(n) for n in range(110,150)]
    original=deepcopy(rows);v=candidate_views(rows)
    assert len(v['candidates'])==100 and len(v['top_candidates'])==25
    assert v['top_candidates'][0]['isin']==rows[110]['isin']
    assert v['top_candidates'][-1]['overall_rank']==25
    assert v['candidate_summary']['basic_qualified_isins']==40 and rows==original
    assert v['top_candidates'][0]['category_size']==150


def test_duplicate_isin_prefers_qualified_listing_and_preserves_category_rank():
    v=candidate_views([row(1,blocks=['SPREAD']),row(1,id='NO0000000001:4',score=900,rank=8),row(2,blocks=['OVERLAPP'])])
    assert len(v['top_candidates'])==2
    assert v['top_candidates'][1]['id'].endswith(':4')
    assert v['top_candidates'][1]['rank']==8 and v['top_candidates'][1]['overall_rank']==2
    assert v['candidate_summary']['basic_qualified_isins']==2


def test_no_trend_no_top_and_stable_ties():
    assert not candidate_views([row(1,score=None)])['top_candidates']
    assert [r['id'] for r in candidate_views([row(2,score=1),row(1,score=1)])['top_candidates']]==[row(1)['id'],row(2)['id']]


def test_country_not_inferred_from_isin_and_fallback_currency_disclosed():
    info=dict(identity_lines(row(1,returns_currency='EUR',returns_currency_source='INSTRUMENT_CURRENCY_FALLBACK')))
    assert info['Registreringsland']=='Ikke oppgitt av kilden'
    assert 'ikke bekreftet' in info['Avkastningsvaluta']
    assert dict(identity_lines(row(1,domicile='Irland',investment_area='Brasil')))['Investeringsland / region']=='Brasil'


def test_cards_escape_untrusted_name_and_do_not_show_none():
    html=candidate_card(row(1,name='<script>alert(1)</script>',risk=None))
    assert '<script>' not in html and '&lt;script&gt;' in html
    assert 'None' not in html and 'Ikke oppgitt' in html
    assert 'Kategori-rang' in html and 'Poeng' in html


def test_pending_identity_outside_saved_hundred_and_date_timezone():
    order={'id':'outside','side':'BUY'}
    s={'fund_identity':{'outside':{'name':'Langt fondsnavn','returns_currency':'USD','price_at':'2026-10-10T12:00:00Z'}}}
    assert resolve_identity(s,order)['name']=='Langt fondsnavn'
    assert date_label('2026-10-10T12:00:00Z')=='10.10.2026 12:00 UTC'
    assert date_label(None)=='Ikke oppgitt'


def test_legacy_snapshot_fallback_and_excel_top25_same_selection():
    from superfund_reports import xlsx_bytes,pdf_bytes
    from openpyxl import load_workbook
    from pypdf import PdfReader
    s={'at':'2026-10-10',**candidate_views([row(n) for n in range(30)])}
    wb=load_workbook(BytesIO(xlsx_bytes(s)))
    assert wb['Topp 25'].max_row==26 and wb['Topp 25']['B2'].value=='Fond 0'
    text=''.join(p.extract_text() for p in PdfReader(BytesIO(pdf_bytes(s))).pages)
    assert 'Fond 24' in text and 'Fond 25' not in text and 'Registreringsland' in text
    assert len(top_candidates({'candidates':[row(1)]}))==1


def test_all_menus_render_long_text_nested_capacity_without_network(monkeypatch):
    from streamlit.testing.v1 import AppTest
    import pages.superfund as page
    s={'at':'2026-10-10T12:00:00Z',**candidate_views([row(1,name='Et svært langt fondsnavn '*10)]),
       'model':{'cash':1000000,'initial_capital':1000000,'positions':{},'orders':[{'id':row(1)['id'],'side':'BUY','reason':'Positiv trend','requested_at':'2026-10-10T12:00:00Z'}]},
       'learning':{},'shadows':{}}
    monkeypatch.setattr(page,'snapshot',lambda:s)
    monkeypatch.setattr(page,'read',lambda name,default:{} if name!='job.json' else {'capacity':{'headroom_mb':1725,'cpu_pressure_avg10_pct':2}})
    monkeypatch.setattr(page,'request_scan',lambda:pytest.fail('Unexpected scan'))
    app=AppTest.from_string('import streamlit as st\nfrom pages.superfund import render_superfund\nrender_superfund(st)').run()
    assert not app.exception
    assert len(app.expander)>=7
    markdown=' '.join(x.value for x in app.markdown)
    assert '1725' in markdown and 'Et svært langt fondsnavn' in markdown
    assert 'max-height:none' in markdown


def test_optional_source_identity_fields_do_not_crash_catalogue():
    from superfund_presentation import identity
    empty=identity(row(1,detail={'company':None,'regions':None}))
    assert 'manager_country' not in empty and 'investment_area' not in empty
    actual=identity(row(1,detail={'company':{'physicalAddress':{'country':'Germany'}},'regions':[{'exposureType':'Latin America','weight':98}]}))
    assert actual['manager_country']=='Germany' and 'Latin America' in actual['investment_area']
    assert 'domicile' not in actual  # manager's address is not fund domicile
