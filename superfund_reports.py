"""One frozen snapshot powers screen, PDF, CSV and workbook exports."""
import csv
from html import escape
from io import BytesIO, StringIO
import json
from superfund_presentation import country_candidates, top_candidates, resolve_identity, identity_lines, RANKING_NOTE, date_label
from superfund_owned_rankings import owned_rankings


def safe_text(value):
    value=str(value)
    return "'"+value if value.startswith(('=','+','-','@')) else value


def csv_bytes(snapshot):
    stream=StringIO(); writer=csv.writer(stream)
    writer.writerow(['ISIN','Fond','Kategori','Kjøpstid','Rang ved kjøp','Rang nå','Poeng ved kjøp','Poeng nå','Rang blant eide','Antall','Verdi NOK','Gevinst NOK','Investeringsland / region','Registreringsland','Avkastningsvaluta'])
    owned={r['id']:r for r in owned_rankings(snapshot)}
    for key,p in snapshot.get('model',{}).get('positions',{}).items():
        current=owned.get(key,{})
        value=p['quantity']*p['last_nok']
        ident=dict(identity_lines(resolve_identity(snapshot,p)))
        writer.writerow([p['isin'],safe_text(p['name']),safe_text(p['category']),p['bought_at'],p['rank_at_buy'],p['rank_now'],p.get('score_at_buy'),current.get('score_now'),current.get('owned_rank'),p['quantity'],value,value-p['cost_nok'],safe_text(ident['Investeringsland / region']),safe_text(ident['Registreringsland']),safe_text(ident['Avkastningsvaluta'])])
    return ('\ufeff'+stream.getvalue()).encode('utf-8')


def xlsx_bytes(snapshot):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title='Superfond'
    for index,row in enumerate(csv.reader(StringIO(csv_bytes(snapshot).decode('utf-8-sig')))):
        if index:
            row[4:6]=[int(x) if x else None for x in row[4:6]]
            row[6:9]=[float(x) if x and i!=8 else int(x) if x else None for i,x in enumerate(row[6:9],6)]
            row[9:12]=[float(x) if x else None for x in row[9:12]]
        ws.append(row)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(45,max(len(str(c.value or '')) for c in col)+2)
    top=wb.create_sheet('Topp 25')
    top.append(['Plass','Fond','ISIN','Type','Investeringsland / region','Registreringsland','Kategori','Avkastningsvaluta','Poeng','Uke %','Måned %','Kvartal %','Risiko /7','Kostnad %','Kursdato','Status'])
    for row in top_candidates(snapshot):
        ident=dict(identity_lines(row));r=row.get('returns',{})
        values=[row.get('overall_rank'),row.get('name'),row.get('isin'),ident['Type'],ident['Investeringsland / region'],ident['Registreringsland'],row.get('category'),ident['Avkastningsvaluta'],row.get('score'),r.get('yield_1w'),r.get('yield_1m'),r.get('yield_3m'),row.get('risk'),row.get('cost_pct'),date_label(row.get('price_at')),' / '.join(row.get('blocks',[])) or 'Grunnkrav oppfylt']
        top.append([safe_text(v) if isinstance(v,str) else v for v in values])
    top.freeze_panes='A2';top.auto_filter.ref=top.dimensions
    for col in top.columns:top.column_dimensions[col[0].column_letter].width=min(55,max(len(str(c.value or '')) for c in col)+2)
    countries=wb.create_sheet('Beste per område')
    countries.append(['Investeringsområde','Plass i området','Fond','ISIN','Kategori','Avkastningsvaluta','Uke %','Måned %','Kvartal %','Risiko /7','Kostnad %','Kursdato'])
    for area,rows in sorted(country_candidates(snapshot).items()):
        for row in rows:
            r=row.get('returns') or {}
            values=[area,row['country_rank'],row.get('name'),row.get('isin'),row.get('category'),dict(identity_lines(row))['Avkastningsvaluta'],r.get('yield_1w'),r.get('yield_1m'),r.get('yield_3m'),row.get('risk'),row.get('cost_pct'),date_label(row.get('price_at'))]
            countries.append([safe_text(v) if isinstance(v,str) else v for v in values])
    countries.freeze_panes='A2';countries.auto_filter.ref=countries.dimensions
    for col in countries.columns:countries.column_dimensions[col[0].column_letter].width=min(55,max(len(str(c.value or '')) for c in col)+2)
    risk=wb.create_sheet('Hendelsesrisiko')
    for key,value in snapshot.get('risk',{}).items():
        risk.append([key,safe_text(json.dumps(value,ensure_ascii=False))])
    risk.column_dimensions['A'].width=32; risk.column_dimensions['B'].width=90
    note=wb.create_sheet('Les først');note.append([RANKING_NOTE]);note.append(['Registreringsland utledes ikke fra ISIN. Ukjente land og valutaopplysninger merkes.'])
    out=BytesIO();wb.save(out);return out.getvalue()


def pdf_bytes(snapshot):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    out=BytesIO();styles=getSampleStyleSheet();parts=[]
    fonts=Path(__file__).resolve().parent/'assets/fonts'
    for name,file in [('SuperfundNoto','NotoSans-Regular.ttf'),('SuperfundNotoBold','NotoSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
    for style in styles.byName.values():
        if not hasattr(style,'fontSize'):continue
        style.fontName='SuperfundNotoBold' if style.name in ('Title','Heading1','Heading2','Heading3') else 'SuperfundNoto'
        style.leading=max(style.leading,style.fontSize*1.4)
    def line(value,style='BodyText'):
        parts.append(Paragraph(escape(str(value)),styles[style]));parts.append(Spacer(1,7))
    line('Superfondportefølje – modell og læring','Title')
    line('Snapshot: '+snapshot.get('at',''))
    line('Programversjon: '+snapshot.get('version','Ikke registrert'))
    model=snapshot.get('model',{});equity=model.get('cash',0)+sum(p['quantity']*p['last_nok'] for p in model.get('positions',{}).values())
    line(f"Verdi NOK {equity:,.2f} · kontanter {model.get('cash',0):,.2f} · største fall {model.get('max_drawdown_pct',0):.2f}%")
    line('MODELLHANDEL. Ingen ordre sendes til Nordnet. Utførelse er et estimat på senere observert kurs; NAV, valuta og ordrefrister kan gi avvik.')
    if snapshot.get('risk'):
        r=snapshot['risk'];line('Hendelsesrisiko','Heading2')
        line(f"Brasil konservativ øvre eksponering {r['brazil_upper_pct']:.1f}% · tak {r['cap_pct']}% · risikoindikator {r['risk_indicator']}/100")
        line(r['indicator_method']);line(r['currency_note'])
        line(f"Scenario: aksjer -10% og valuta -10%, samlet -19%. Tap ved øvre eksponering NOK {r['stress']['portfolio_loss_upper_nok']:,.0f}. Ikke prognose.")
        for proposal in r['proposals']:line(f"Krever godkjenning: {proposal['name']} · reduser {proposal['fraction']*100:.1f}%")
    line('Eide fond – aktuelle poeng og avkastning','Heading2')
    for item in sorted(owned_rankings(snapshot),key=lambda x:(x['score_now'] is None,-(x['score_now'] or 0),x['name'])):
        points=f"{item['score_now']:.3f}" if item['score_now'] is not None else 'mangler grunnlag'
        profit=f"{item['return_pct']:+.2f}%" if item['return_pct'] is not None else 'ukjent'
        line(f"{item['owned_rank'] or '—'}. {item['name']} · poeng nå {points} · ved kjøp {item['score_at_buy']} · siden kjøp {profit} · kursdato {date_label(item['price_at'])}")
    line('Beholdninger','Heading2')
    for p in model.get('positions',{}).values():
        line(p['name']+' · '+p['isin'],'Heading3')
        for label,value in identity_lines(resolve_identity(snapshot,p)):line(f'{label}: {value}')
        line(f"Kjøp NOK {p['entry_nok']:.2f} · nå {p['last_nok']:.2f} · topp {p['peak_nok']:.2f}")
        line(f"Rang ved kjøp {p['rank_at_buy']} → nå {p['rank_now']} · verdi {p['quantity']*p['last_nok']:,.2f}")
        if p.get('floor_nok'):line(f"Salgsutløser NOK {p['floor_nok']:.2f}. Utførelseskurs er ikke garantert.")
    line('Ventende modellordre','Heading2')
    for o in model.get('orders',[]):
        ident=resolve_identity(snapshot,o)
        line(f"{o['side']} · {ident.get('name',o['id'])} · {o['reason']}")
        for label,value in identity_lines(ident):line(f'{label}: {value}')
        line('Opprettet: '+date_label(o.get('requested_at'))+' · tidligst behandling: '+date_label(o.get('next_at')))
        line('Krever en senere observert kursdato og ny kontroll av kjøpskrav.')
    line('Topp 25 – grunnkrav oppfylt','Heading2')
    line(RANKING_NOTE)
    if 'top_candidates' not in snapshot:line('Eldre snapshot: utvalget er begrenset til de 100 lagrede kandidatene.')
    for row in top_candidates(snapshot):
        start=len(parts)
        line(f"{row.get('overall_rank','—')}. {row['name']} · {row['isin']}",'Heading3')
        for label,value in identity_lines(row):line(f'{label}: {value}')
        r=row.get('returns',{})
        line(f"Kategori-rang {row.get('rank')} av {row.get('category_size','ukjent')} · poeng {row.get('score')} · uke {r.get('yield_1w')}% · måned {r.get('yield_1m')}% · kvartal {r.get('yield_3m')}%")
        line('Kursdato: '+date_label(row.get('price_at'))+' · status: '+(' / '.join(row.get('blocks',[])) or 'Grunnkrav oppfylt'))
        chunk=parts[start:];del parts[start:];parts.append(KeepTogether(chunk))
    line('Datadekning og diagnoser','Heading2')
    for key,value in snapshot.get('coverage',{}).items():line(f'{key}: {value}')
    line('Læring og skyggetester','Heading2')
    for name,state in snapshot.get('shadows',{}).items():
        value=state.get('cash',0)+sum(p['quantity']*p['last_nok'] for p in state.get('positions',{}).values())
        line(f"{name}: NOK {value:,.2f} · fall {state.get('max_drawdown_pct',0):.2f}% · {len(state.get('history',[]))} observasjoner")
    line('Korte serier dokumenterer ikke meravkastning. Skyggeregler er fryst ved oppstart og endrer ikke hovedporteføljen.')
    line('Kostnader: modellfriksjon i basispunkter pluss halv ETF-spread per handel og valutapåslag. NAV inneholder fondets løpende kostnader; Nordnets samlede kostnad er et utvalgsfilter og trekkes ikke dobbelt. Plattformkostnader/refusjoner er ennå ikke avstemt mot faktisk konto.')
    line('Regler: '+json.dumps(snapshot.get('parameters',{}),ensure_ascii=False))
    SimpleDocTemplate(out,pagesize=A4,rightMargin=36,leftMargin=36).build(parts)
    return out.getvalue()


def prepare_report_link(snapshot):
    """Authenticated application route; never publish a public PDF payload.

    Downloads are generated from the complete stored model snapshot after the
    application's normal login check. No external report storage or new
    anonymous access token is created.
    """
    import os
    from urllib.parse import urlsplit
    base=(os.getenv('RENDER_EXTERNAL_URL') or 'https://aksje-app.onrender.com').rstrip('/')
    parsed=urlsplit(base)
    if parsed.scheme!='https' or not parsed.netloc:raise ValueError('App-lenken må være HTTPS')
    return {'url':base+'/?aa_nav=superfund','at':snapshot['at'],
            'run_id':'SF-'+snapshot.get('model',{}).get('last_frame','unknown'),
            'access':'APP_LOGIN_REQUIRED','delivery':'APP_DOWNLOADS'}

def _fund_chart_flowable(points):
    """Draw observed NOK quote points only; no synthetic interpolation of missing dates."""
    from datetime import datetime
    from math import isfinite
    from reportlab.graphics.shapes import Drawing, PolyLine, String, Line
    from reportlab.lib.colors import HexColor
    values={}
    for obs in points or []:
        try:
            stamp=str(obs.get('price_at') or '')
            dt=datetime.fromisoformat(stamp.replace('Z','+00:00'))
            price=float(obs.get('nok_price'))
            if price>0 and isfinite(price) and dt.tzinfo is not None:
                values[dt]=price
        except (TypeError, ValueError, OverflowError):
            continue
    ordered=sorted(values.items())
    if len(ordered)<2:
        return None
    low=min(v for _,v in ordered);high=max(v for _,v in ordered)
    width,height=475,170
    plot_left,plot_right,plot_low,plot_high=48,460,35,145
    span=max(1.0,(ordered[-1][0]-ordered[0][0]).total_seconds())
    value_span=high-low or max(1.0,abs(high)*0.01)
    coords=[]
    for dt,price in ordered:
        x=plot_left+(dt-ordered[0][0]).total_seconds()/span*(plot_right-plot_left)
        y=plot_low+(price-low)/value_span*(plot_high-plot_low)
        coords.extend([x,y])
    chart=Drawing(width,height)
    chart.add(Line(plot_left,plot_low,plot_right,plot_low,strokeColor=HexColor('#6b7280')))
    chart.add(PolyLine(coords,strokeColor=HexColor('#1768a8'),strokeWidth=2))
    chart.add(String(2,plot_high,f'{high:,.2f} NOK',fontSize=8))
    chart.add(String(2,plot_low,f'{low:,.2f} NOK',fontSize=8))
    chart.add(String(plot_left,12,ordered[0][0].strftime('%d.%m.%Y'),fontSize=8))
    chart.add(String(plot_right-62,12,ordered[-1][0].strftime('%d.%m.%Y'),fontSize=8))
    return chart


def fund_pdf_bytes(snapshot, position_id):
    """Print-ready per-fund report; no navigation controls or external sharing token."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    positions=snapshot.get('model',{}).get('positions',{})
    pos=positions[position_id]
    ranking=next((r for r in owned_rankings(snapshot) if r['id']==position_id),{})
    from pathlib import Path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    styles=getSampleStyleSheet();parts=[];output=BytesIO()
    fonts=Path(__file__).resolve().parent/'assets/fonts'
    for font_name,filename in [('SuperfundNoto','NotoSans-Regular.ttf'),
                               ('SuperfundNotoBold','NotoSans-Bold.ttf')]:
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name,str(fonts/filename)))
    for style in styles.byName.values():
        if hasattr(style,'fontName'):
            style.fontName='SuperfundNotoBold' if style.name in ('Title','Heading1','Heading2','Heading3') else 'SuperfundNoto'
            style.leading=max(style.leading,style.fontSize*1.4)
    def write(value,style='Normal'):
        parts.append(Paragraph(escape(str(value)),styles[style]))
        parts.append(Spacer(1,8))
    write(pos['name'],'Title')
    write('Superfond modellrapport | '+str(snapshot.get('at','')))
    write('ISIN: '+str(pos.get('isin','')))
    write('Modellportefølje. Ikke dokumentasjon av en reell meglerordre.')
    score=ranking.get('score_now')
    write('Oppdatert poeng: '+(f'{score:.3f}' if score is not None else 'Utilgjengelig – manglende trenddata'))
    write('Rang blant eide fond: '+str(ranking.get('owned_rank') or 'ikke tilgjengelig'))
    write('Poeng ved kjøp: '+str(pos.get('score_at_buy','ukjent')))
    write('Resultat siden kjøp: '+(f"{ranking['return_pct']:+.2f}%" if ranking.get('return_pct') is not None else 'ukjent'))
    write('Siste modellkurs (NOK): '+str(pos.get('last_nok')))
    write('Trailing-stop-grense (NOK): '+str(pos.get('floor_nok')))
    risk_row=next((r for r in snapshot.get('risk',{}).get('positions',[]) if r.get('id')==position_id),None)
    if risk_row:
        write('Risikostatus: '+str(risk_row.get('stop_status') or 'Ikke tilgjengelig'))
        drop=risk_row.get('drop_from_peak_pct')
        write('Fall fra topp: '+(f'{float(drop):.2f}%' if drop is not None else 'ikke tilgjengelig'))
        write('Risikogrunnlag: '+str(risk_row.get('source') or 'Ikke tilgjengelig'))
    else:
        write('Risikostatus: Ikke tilgjengelig i rapportens snapshot')
    ident=resolve_identity(snapshot,pos)
    write('Kursutvikling · daterte observasjoner i NOK','Heading2')
    chart=_fund_chart_flowable(ident.get('price_history_nok') or [])
    if chart is not None:
        parts.append(chart)
        parts.append(Spacer(1,10))
        write('Faktiske observasjonstidspunkter, ikke utfylte handelsdager. Historikk kan være kortere enn valgt grafperiode.')
    else:
        write('Kursgraf utilgjengelig: færre enn to gyldige daterte NAV-/kursobservasjoner.')
    for label,value in identity_lines(ident):write(label+': '+str(value))
    for title,key in [('Land og regioner','regions'),('Sektorer','sectors'),('Største beholdninger','holdings')]:
        write(title,'Heading2')
        rows=pos.get(key) or []
        if not rows:write('Ikke tilgjengelig i siste rapporterte datasett.')
        for item in rows[:30]:
            write(str(item.get('displayName') or item.get('name') or item.get('exposureType') or 'Ukjent')+' · '+str(item.get('weight','?'))+' %')
    write('Fordeling og selskaper er rapporterte fondsdata, ikke sanntidsbeholdning. Rapportert datotid: '+str(pos.get('regions_verified_at') or 'ukjent'))
    SimpleDocTemplate(output,pagesize=A4).build(parts)
    return output.getvalue()
